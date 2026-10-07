from __future__ import annotations

import io
import re
from dataclasses import dataclass

import fitz
from docx import Document


class ExtractionError(ValueError):
    """Raised when an uploaded document cannot produce usable text."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    filename: str
    content_type: str

    @property
    def character_count(self) -> int:
        return len(self.text)


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(character for character in text if character in "\n\t" or ord(character) >= 32)
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return text.strip()


def _ensure_text(text: str, filename: str) -> str:
    normalized = normalize_text(text)
    if not normalized:
        raise ExtractionError(f"Could not extract text from {filename}. The document is empty or contains no extractable text.")
    return normalized


def extract_pdf_text(content: bytes, filename: str = "document.pdf") -> str:
    try:
        with fitz.open(stream=content, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document)
    except Exception as exc:
        raise ExtractionError(f"Could not read PDF file {filename}.") from exc
    return _ensure_text(text, filename)


def extract_docx_text(content: bytes, filename: str = "document.docx") -> str:
    try:
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
        text = "\n".join(paragraphs)
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX file {filename}.") from exc
    return _ensure_text(text, filename)


def extract_txt_text(content: bytes, filename: str = "document.txt") -> str:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError(f"Could not read text file {filename}. Use UTF-8 encoding.") from exc
    return _ensure_text(text, filename)


def extract_file_text(
    content: bytes,
    filename: str,
    content_type: str = "application/octet-stream",
) -> ExtractedDocument:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    extractors = {
        "pdf": extract_pdf_text,
        "docx": extract_docx_text,
        "txt": extract_txt_text,
    }
    extractor = extractors.get(extension)
    if extractor is None:
        raise ExtractionError(f"Unsupported file type for {filename}.")
    return ExtractedDocument(
        text=extractor(content, filename),
        filename=filename,
        content_type=content_type,
    )
